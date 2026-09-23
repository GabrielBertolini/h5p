import base64
import io
import json
import os
import time
import uuid
import zipfile

import markdown
import streamlit as st
from google import genai
from google.genai import types

# Tenta carregar o dotenv localmente; no Streamlit Cloud o segredo virá de st.secrets
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

# Configuração da página (layout centralizado)
st.set_page_config(page_title="Gerador de Accordion H5P", layout="centered")

# Recupera chave da API (localmente do .env ou das Secrets do Streamlit Cloud)
api_key = os.getenv("GEMINI_API_KEY") or st.secrets.get("GEMINI_API_KEY")

TXT_BASE64_PATH = "accordion.txt"

MODELOS_PARA_TENTAR = [
    "gemini-3.5-flash-lite",
    "gemini-3.8-flash",
    "gemini-3.6-flash",
]

# Anexos
TIPOS_ANEXO = ["pdf", "txt", "md", "csv", "docx", "png", "jpg", "jpeg", "webp"]
MIME_BINARIOS = {
    "pdf": "application/pdf",
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "webp": "image/webp",
}
LIMITE_ANEXOS_MB = 18  # limite de envio direto para a API do Gemini (~20 MB)


class ErroAmigavel(Exception):
    """Erro cuja mensagem já está pronta para ser mostrada ao usuário."""


# ---------------------------------------------------------------------------
# Funções auxiliares
# ---------------------------------------------------------------------------
def obter_molde_base64() -> str:
    if os.path.exists(TXT_BASE64_PATH):
        with open(TXT_BASE64_PATH, "r", encoding="utf-8") as f:
            return f.read().strip()
    return ""


def novo_painel(titulo: str = "", conteudo: str = "") -> str:
    """Cria um painel no estado da sessão e devolve seu id."""
    pid = uuid.uuid4().hex[:8]
    st.session_state[f"tit_{pid}"] = titulo
    st.session_state[f"cont_{pid}"] = conteudo
    return pid


def decodificar_texto(dados: bytes) -> str:
    for codificacao in ("utf-8", "latin-1"):
        try:
            return dados.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return dados.decode("utf-8", errors="replace")


def extrair_docx(dados: bytes, nome: str) -> str:
    try:
        import docx  # pacote python-docx
    except ImportError:
        raise ErroAmigavel(
            f"Não consegui ler o arquivo **{nome}**: o suporte a Word (.docx) "
            "não está instalado. Envie o arquivo em PDF ou adicione "
            "`python-docx` ao requirements.txt."
        )
    documento = docx.Document(io.BytesIO(dados))
    linhas = [p.text for p in documento.paragraphs if p.text.strip()]
    for tabela in documento.tables:
        for linha in tabela.rows:
            linhas.append(" | ".join(c.text.strip() for c in linha.cells))
    return "\n".join(linhas)


def preparar_anexos(arquivos):
    """Devolve (partes_binarias, textos) prontos para enviar ao Gemini."""
    partes, textos = [], []
    total = sum(len(a.getvalue()) for a in arquivos or [])
    if total > LIMITE_ANEXOS_MB * 1024 * 1024:
        raise ErroAmigavel(
            f"Os anexos somam {total / 1024 / 1024:.1f} MB, e o limite é "
            f"{LIMITE_ANEXOS_MB} MB. Tente enviar menos arquivos ou versões menores."
        )

    for arq in arquivos or []:
        dados = arq.getvalue()
        ext = arq.name.rsplit(".", 1)[-1].lower()
        try:
            if ext in MIME_BINARIOS:
                partes.append(
                    types.Part.from_bytes(data=dados, mime_type=MIME_BINARIOS[ext])
                )
            elif ext == "docx":
                textos.append((arq.name, extrair_docx(dados, arq.name)))
            else:
                textos.append((arq.name, decodificar_texto(dados)))
        except ErroAmigavel:
            raise
        except Exception:
            raise ErroAmigavel(
                f"Não consegui ler o arquivo **{arq.name}**. Ele pode estar "
                "corrompido ou protegido por senha. Tente removê-lo ou enviar em outro formato."
            )
    return partes, textos


def erro_temporario(err: Exception) -> bool:
    txt = str(err).lower()
    return any(
        s in txt
        for s in ("503", "unavailable", "overloaded", "high demand",
                  "500", "internal", "429", "resource_exhausted")
    )


def chamar_gemini(client, contents) -> str:
    """Tenta os modelos em sequência; se todos estiverem sobrecarregados,
    espera um pouco e tenta mais uma rodada."""
    ultimo_erro = None
    for rodada in range(2):
        for mod in MODELOS_PARA_TENTAR:
            try:
                response = client.models.generate_content(
                    model=mod,
                    contents=contents,
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json"
                    ),
                )
                if response and response.text:
                    return response.text
            except Exception as err:
                ultimo_erro = err
        if rodada == 0 and ultimo_erro is not None and erro_temporario(ultimo_erro):
            time.sleep(4)
        else:
            break
    raise ultimo_erro or RuntimeError("A IA devolveu uma resposta vazia.")


def gerar_paineis_ia(tema, num_paineis, extensao, instrucao, texto_fonte, anexos) -> list:
    partes_anexos, textos_anexos = preparar_anexos(anexos)

    bloco_anexos = ""
    if textos_anexos:
        bloco_anexos += "\n\nConteúdo dos anexos de texto:\n"
        for nome, texto in textos_anexos:
            bloco_anexos += f"\n### {nome}\n{texto}\n"
    if partes_anexos:
        bloco_anexos += (
            "\n\nTambém foram anexados arquivos (PDFs/imagens) junto desta mensagem. "
            "Use-os como material de apoio."
        )

    prompt = f"""Você é um assistente pedagógico. Crie um conteúdo para a ferramenta Accordion do H5P.
Tema: {tema}
Número de painéis: {num_paineis}
Extensão aproximada: {extensao}
Instruções adicionais: {instrucao}
Texto base: {texto_fonte}{bloco_anexos}

Gere uma lista JSON válida contendo os painéis.
O campo "content" deve estar em Markdown simples (parágrafos separados por linha em branco,
**negrito**, *itálico* e listas com "-"). Não use HTML.
Formato obrigatório de cada item:
[
  {{
    "title": "Título do Painel",
    "content": "Texto do painel em Markdown."
  }}
]"""

    client = genai.Client(api_key=api_key)
    texto_resposta = chamar_gemini(client, [prompt, *partes_anexos])

    dados = json.loads(texto_resposta)
    # Às vezes o modelo devolve {"panels": [...]} em vez da lista direta
    if isinstance(dados, dict):
        dados = next((v for v in dados.values() if isinstance(v, list)), [])
    if not dados:
        raise ErroAmigavel(
            "A IA não devolveu nenhum painel desta vez. Tente gerar de novo, "
            "talvez com uma instrução um pouco mais detalhada."
        )
    return dados


def mensagem_amigavel(err: Exception) -> str:
    """Traduz erros técnicos em mensagens humanas."""
    if isinstance(err, ErroAmigavel):
        return str(err)
    if isinstance(err, json.JSONDecodeError):
        return (
            "A IA respondeu num formato inesperado. Clique em **Gerar accordion** "
            "de novo; normalmente funciona na segunda tentativa."
        )

    txt = str(err).lower()
    if any(s in txt for s in ("503", "unavailable", "overloaded", "high demand")):
        return (
            "A IA está com muita demanda neste momento e não conseguiu responder. "
            "Isso costuma passar rápido: aguarde alguns segundos e tente de novo."
        )
    if any(s in txt for s in ("429", "resource_exhausted", "quota", "rate limit")):
        return "Atingimos o limite de uso da IA por agora. Espere um minuto e tente novamente."
    if any(s in txt for s in ("api key", "api_key", "401", "403", "permission", "unauthenticated")):
        return (
            "Não consegui me conectar à IA porque a chave de acesso parece inválida. "
            "Confira a GEMINI_API_KEY nas configurações do app."
        )
    if "404" in txt or "not found" in txt:
        return (
            "Os modelos de IA configurados não estão disponíveis no momento. "
            "Verifique os nomes dos modelos no código."
        )
    if any(s in txt for s in ("too large", "exceeds", "payload", "token count")):
        return (
            "O material enviado é grande demais para a IA processar de uma vez. "
            "Tente reduzir o texto fonte ou enviar menos anexos."
        )
    if any(s in txt for s in ("timeout", "timed out", "connection", "network")):
        return "Não consegui me conectar à IA. Verifique a conexão e tente novamente."
    return "Algo deu errado ao gerar o conteúdo. Tente novamente em instantes."


def md_para_html(texto: str) -> str:
    return markdown.markdown(texto or "", extensions=["extra", "nl2br", "sane_lists"])


def montar_h5p(titulo_geral: str, paineis: list) -> bytes:
    """paineis: lista de dicts {"title": str, "content": str (markdown)}."""
    b64_string = obter_molde_base64()
    if not b64_string:
        raise ErroAmigavel(
            f"Não encontrei o molde do accordion ('{TXT_BASE64_PATH}'). "
            "Verifique se o arquivo está no repositório."
        )
    molde_buffer = io.BytesIO(base64.b64decode(b64_string))

    content_h5p = {
        "panels": [
            {
                "title": p["title"] or f"Painel {i+1}",
                "content": {
                    "params": {"text": md_para_html(p["content"])},
                    "library": "H5P.AdvancedText 1.1",
                    "subContentId": str(uuid.uuid4()),
                    "metadata": {
                        "contentType": "Text",
                        "license": "U",
                        "title": (p["title"] or "Untitled Text")[:60],
                    },
                },
            }
            for i, p in enumerate(paineis)
        ],
        "hTag": "h2",
    }

    output_buffer = io.BytesIO()
    with zipfile.ZipFile(molde_buffer, "r") as zip_in:
        with zipfile.ZipFile(output_buffer, "w", zipfile.ZIP_DEFLATED) as zip_out:
            for item in zip_in.infolist():
                if item.filename == "content/content.json":
                    zip_out.writestr(
                        "content/content.json",
                        json.dumps(content_h5p, ensure_ascii=False, indent=2),
                    )
                elif item.filename == "h5p.json":
                    try:
                        h5p_data = json.loads(zip_in.read(item.filename).decode("utf-8"))
                        h5p_data["title"] = titulo_geral
                        zip_out.writestr(
                            "h5p.json",
                            json.dumps(h5p_data, ensure_ascii=False, indent=2),
                        )
                    except Exception:
                        zip_out.writestr(item.filename, zip_in.read(item.filename))
                else:
                    zip_out.writestr(item.filename, zip_in.read(item.filename))

    return output_buffer.getvalue()


def coletar_dados():
    """Lê o título geral e os painéis atuais do editor."""
    titulo_geral = (st.session_state.get("titulo_geral") or "").strip()
    dados = [
        {
            "title": (st.session_state.get(f"tit_{pid}") or "").strip(),
            "content": st.session_state.get(f"cont_{pid}") or "",
        }
        for pid in st.session_state.get("paineis", [])
    ]
    return titulo_geral, dados


# ---------------------------------------------------------------------------
# Callbacks do editor (rodam antes do próximo render)
# ---------------------------------------------------------------------------
def adicionar_painel():
    st.session_state.paineis.append(novo_painel("Novo painel", ""))


def remover_painel(pid):
    st.session_state.paineis.remove(pid)


def mover_painel(pid, delta):
    lista = st.session_state.paineis
    i = lista.index(pid)
    j = i + delta
    if 0 <= j < len(lista):
        lista[i], lista[j] = lista[j], lista[i]


# ---------------------------------------------------------------------------
# Interface: formulário
# ---------------------------------------------------------------------------
st.title("Gerador de Accordion H5P", anchor=False)
st.divider()

st.caption("CONTEÚDO")

tema = st.text_input("Tema *", placeholder="ex.: Fotossíntese")

instrucao = st.text_area(
    "Instrução",
    help="Foco, público, o que priorizar ou evitar.",
    placeholder="ex.: Material de revisão para alunos do 1º ano do EM...",
)

texto_fonte = st.text_area(
    "Texto fonte",
    help="Adicionar material de apoio gera resultados mais personalizados.",
    placeholder="Cole aqui o texto para apoio...",
)

anexos = st.file_uploader(
    "Anexos (opcional)",
    type=TIPOS_ANEXO,
    accept_multiple_files=True,
    help="PDFs, textos, documentos Word ou imagens para complementar o conteúdo. "
    f"Até {LIMITE_ANEXOS_MB} MB no total.",
)

col1, col2 = st.columns(2)
with col1:
    num_paineis = st.number_input("Painéis", min_value=1, max_value=15, value=5)
with col2:
    extensao = st.selectbox("Extensão", ["Curto", "Médio", "Longo"], index=1)

btn_gerar = st.button("Gerar accordion", type="primary", use_container_width=True)

# Status logo abaixo do botão
status_area = st.empty()
if st.session_state.get("paineis") is not None:
    status_area.success("Pronto! Revise e edite o accordion abaixo.")

# Processamento ao clicar em "Gerar accordion"
if btn_gerar:
    if not tema.strip():
        status_area.warning("Para começar, informe o **tema** do accordion.")
    elif not api_key:
        status_area.error(
            "A chave da IA (GEMINI_API_KEY) não está configurada. "
            "Adicione-a nas Secrets do app para poder gerar conteúdo."
        )
    else:
        try:
            with status_area.container():
                with st.spinner("A IA está processando o conteúdo..."):
                    paineis_json = gerar_paineis_ia(
                        tema, num_paineis, extensao, instrucao, texto_fonte, anexos
                    )
            st.session_state.paineis = [
                novo_painel(p.get("title", f"Painel {i+1}"), p.get("content", ""))
                for i, p in enumerate(paineis_json)
            ]
            st.session_state.titulo_geral = tema
            st.session_state.h5p_final = None
            status_area.success("Pronto! Revise e edite o accordion abaixo.")
        except Exception as e:
            with status_area.container():
                st.error(mensagem_amigavel(e))
                if not isinstance(e, ErroAmigavel):
                    with st.expander("Detalhes técnicos"):
                        st.code(str(e), language=None)

# ---------------------------------------------------------------------------
# Editor
# ---------------------------------------------------------------------------
if st.session_state.get("paineis") is not None:
    st.divider()
    st.subheader("Editar accordion", anchor=False)

    st.text_input("Título geral", key="titulo_geral")
    st.caption(
        "O texto aceita Markdown: **negrito**, *itálico*, listas com `-` "
        "e parágrafos separados por uma linha em branco."
    )

    paineis = st.session_state.paineis
    total = len(paineis)

    for i, pid in enumerate(paineis):
        with st.container(border=True):
            st.text_input(f"Título do painel {i+1}", key=f"tit_{pid}")

            aba_editar, aba_previa = st.tabs(["Editar", "Pré-visualizar"])
            with aba_editar:
                st.text_area(
                    "Conteúdo",
                    key=f"cont_{pid}",
                    height=200,
                    label_visibility="collapsed",
                )
            with aba_previa:
                st.markdown(st.session_state.get(f"cont_{pid}") or "_(vazio)_")

            b1, b2, b3 = st.columns(3)
            b1.button(
                "Subir", key=f"up_{pid}", icon=":material/arrow_upward:",
                on_click=mover_painel, args=(pid, -1),
                disabled=(i == 0), use_container_width=True,
            )
            b2.button(
                "Descer", key=f"down_{pid}", icon=":material/arrow_downward:",
                on_click=mover_painel, args=(pid, 1),
                disabled=(i == total - 1), use_container_width=True,
            )
            b3.button(
                "Remover", key=f"del_{pid}", icon=":material/delete:",
                on_click=remover_painel, args=(pid,),
                disabled=(total == 1), use_container_width=True,
            )

    st.button(
        "Adicionar painel", icon=":material/add:",
        on_click=adicionar_painel, use_container_width=True,
    )

    st.write("")

    titulo_geral, dados = coletar_dados()
    assinatura = json.dumps([titulo_geral, dados], ensure_ascii=False)

    if st.button("Gerar arquivo .h5p final", type="primary", use_container_width=True):
        st.session_state.h5p_final = None
        if not titulo_geral:
            st.warning("Dê um **título geral** ao accordion antes de gerar o arquivo.")
        elif not dados:
            st.warning("Adicione pelo menos um painel.")
        else:
            try:
                nome_sanitizado = "".join(
                    c for c in titulo_geral if c.isalnum() or c in (" ", "_")
                ).strip() or "accordion"
                st.session_state.h5p_final = {
                    "bytes": montar_h5p(titulo_geral, dados),
                    "nome": f"accordion_{nome_sanitizado}.h5p",
                    "assinatura": assinatura,
                }
            except Exception as e:
                st.error(
                    mensagem_amigavel(e) if isinstance(e, ErroAmigavel)
                    else "Não consegui montar o arquivo .h5p. Tente novamente."
                )

    # O botão de download é renderizado em TODA execução enquanto o arquivo
    # existir; assim o Streamlit não descarta o arquivo antes do navegador baixá-lo.
    final = st.session_state.get("h5p_final")
    if final:
        if final["assinatura"] == assinatura:
            st.success("Arquivo final gerado!")
            st.download_button(
                label="Baixar arquivo .h5p",
                data=final["bytes"],
                file_name=final["nome"],
                mime="application/zip",
                type="primary",
                icon=":material/download:",
                on_click="ignore",
                use_container_width=True,
            )
        else:
            st.info(
                "Você alterou o conteúdo depois de gerar o arquivo. "
                "Clique em **Gerar arquivo .h5p final** novamente."
            )
